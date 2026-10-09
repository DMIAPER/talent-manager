import io
import os
import re
from pathlib import Path
from typing import Dict, Any, List, Optional
from docx import Document

ASSETS_DIR = Path(__file__).resolve().parent.parent.parent / "assets"
CV_MAESTRO_PATH = ASSETS_DIR / "cv-maestro.md"

def is_corrupted_or_useless_text(text: str) -> bool:
    """
    Detecta si el texto extraído es basura de glifos no decodificados (/i255/, (cid:), etc.)
    o contiene una cantidad insuficiente de palabras legibles.
    """
    if not text or len(text.strip()) < 20:
        return True
    
    # Detección de glifos crudos de fuentes Identity-H / Type 3 (ej: /i255/1/2/3... o /c12/)
    if re.search(r'/[a-zA-Z]\d+/', text) or "(cid:" in text:
        return True
    
    # Densidad excesiva de barras oblicuas '/' (típico de fuentes mal decodificadas)
    slash_count = text.count('/')
    if slash_count > 10 and (slash_count / max(len(text), 1)) > 0.05:
        return True
    
    # Comprobar si hay al menos algunas palabras reconocibles con caracteres alfabéticos
    words = [w for w in text.split() if len(w) >= 2 and any(c.isalpha() for c in w)]
    if len(words) < 5:
        return True
    
    return False

def extract_text_with_pymupdf(file_bytes: bytes) -> str:
    """Motor 1: PyMuPDF (fitz) - Motor C con soporte avanzado de fuentes y tablas CMap."""
    try:
        import pymupdf
        doc = pymupdf.open(stream=file_bytes, filetype="pdf")
        pages_text = []
        for page_idx, page in enumerate(doc):
            # sort=True ordena bloques de texto geométricamente (crucial para diplomas apaisados/horizontales)
            txt = page.get_text("text", sort=True) or ""
            if txt.strip():
                pages_text.append(f"\n--- Página {page_idx + 1} ---\n{txt.strip()}")
        doc.close()
        return "\n".join(pages_text).strip()
    except Exception as e:
        print(f"[PyMuPDF] Falló extracción: {e}")
        return ""

def extract_text_with_pdfplumber(file_bytes: bytes) -> str:
    """Motor 2: pdfplumber - Extractor con análisis de layout y tablas."""
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            pages_text = []
            for page_idx, page in enumerate(pdf.pages):
                txt = page.extract_text(layout=True) or page.extract_text() or ""
                if txt.strip():
                    pages_text.append(f"\n--- Página {page_idx + 1} ---\n{txt.strip()}")
            return "\n".join(pages_text).strip()
    except Exception as e:
        print(f"[pdfplumber] Falló extracción: {e}")
        return ""

def extract_text_with_pypdf(file_bytes: bytes) -> str:
    """Motor 3: pypdf - Extractor estándar de respaldo."""
    try:
        from pypdf import PdfReader
        pdf_reader = PdfReader(io.BytesIO(file_bytes))
        pages_text = []
        for page_idx, page in enumerate(pdf_reader.pages):
            txt = page.extract_text() or ""
            if txt.strip():
                pages_text.append(f"\n--- Página {page_idx + 1} ---\n{txt.strip()}")
        return "\n".join(pages_text).strip()
    except Exception as e:
        print(f"[pypdf] Falló extracción: {e}")
        return ""

def extract_text_with_gemini_vision(file_bytes: bytes) -> str:
    """
    Motor 4: OCR Multimodal Gemini Visión.
    Si todos los extractores locales devuelven glifos corruptos o el PDF es una imagen escaneada,
    Gemini 've' el documento visualmente y extrae el contenido íntegro sin depender de tablas CMap.
    """
    try:
        from backend.services.agent_browser import get_gemini_client
        from google.genai import types
        client, _ = get_gemini_client()
        if not client:
            return ""

        pdf_part = types.Part.from_bytes(data=file_bytes, mime_type="application/pdf")
        prompt = (
            "Extrae y transcribe fielmente todo el texto de este currículum en formato Markdown limpio. "
            "Respeta escrupulosamente los nombres, datos de contacto, enlaces, resumen profesional, "
            "experiencia laboral, formación, fechas, proyectos y habilidades. "
            "No inventes ningún dato que no esté en el documento visual."
        )

        candidate_models = ["gemini-3.5-flash-lite", "gemini-flash-latest"]
        for m in candidate_models:
            try:
                res = client.models.generate_content(
                    model=m,
                    contents=[pdf_part, prompt]
                )
                if res and res.text and not is_corrupted_or_useless_text(res.text):
                    return res.text.strip()
            except Exception as e:
                print(f"[Gemini Vision OCR] Intento con modelo {m} falló: {e}")
                continue
    except Exception as e:
        print(f"[Gemini Vision OCR] Error general: {e}")
    return ""

def extract_text_from_bytes(file_bytes: bytes, filename: str) -> str:
    """
    Extrae texto de un archivo (PDF, DOCX, TXT, MD) aplicando una cascada multicapa
    con detección activa de glifos corruptos.
    """
    ext = filename.lower().split(".")[-1]
    extracted_text = ""

    if ext == "pdf":
        # 1. Intentar PyMuPDF (muy superior en fuentes embebidas y CMaps)
        txt = extract_text_with_pymupdf(file_bytes)
        if txt and not is_corrupted_or_useless_text(txt):
            extracted_text = txt

        # 2. Intentar pdfplumber
        if not extracted_text:
            txt = extract_text_with_pdfplumber(file_bytes)
            if txt and not is_corrupted_or_useless_text(txt):
                extracted_text = txt

        # 3. Intentar pypdf
        if not extracted_text:
            txt = extract_text_with_pypdf(file_bytes)
            if txt and not is_corrupted_or_useless_text(txt):
                extracted_text = txt

        # 4. Fallback visual: Gemini Multimodal Vision OCR
        if not extracted_text:
            print("[CV Parser] Detectado PDF con fuentes corruptas o escaneado. Activando Gemini Vision OCR...")
            txt = extract_text_with_gemini_vision(file_bytes)
            if txt and not is_corrupted_or_useless_text(txt):
                extracted_text = txt

    elif ext in ["docx", "doc"]:
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    paragraphs.append(row_text)
        extracted_text = "\n\n".join(paragraphs)

    elif ext in ["txt", "md"]:
        try:
            extracted_text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            extracted_text = file_bytes.decode("latin-1", errors="ignore")
    else:
        raise ValueError(f"Formato no soportado: .{ext}. Usa PDF, Word (DOCX) o TXT/MD.")

    cleaned_text = extracted_text.strip()
    if not cleaned_text or is_corrupted_or_useless_text(cleaned_text):
        raise ValueError(
            "No se pudo extraer texto legible del archivo PDF. "
            "El documento puede tener tipografías no estándar sin tabla ToUnicode/CID o ser una imagen. "
            "Solución recomendada: ábrelo en tu lector y guárdalo usando 'Imprimir como PDF' o súbelo en formato Word (.docx)."
        )

    return cleaned_text

def save_cv_as_master(cleaned_text: str, filename: str = "") -> Dict[str, Any]:
    ASSETS_DIR.mkdir(parents=True, exist_ok=True)
    
    content = f"# Currículum Vitae (Fuente de Verdad)\n"
    if filename:
        content += f"> Documento original importado: `{filename}`\n\n"
    content += cleaned_text

    with open(CV_MAESTRO_PATH, "w", encoding="utf-8") as f:
        f.write(content)

    return {
        "success": True,
        "path": str(CV_MAESTRO_PATH),
        "char_count": len(cleaned_text),
        "preview": cleaned_text[:300] + ("..." if len(cleaned_text) > 300 else "")
    }

def get_current_cv_info() -> Dict[str, Any]:
    if not CV_MAESTRO_PATH.exists():
        return {
            "exists": False,
            "filename": "cv-maestro.md",
            "content": "",
            "char_count": 0,
            "updated_at": None
        }

    stat = CV_MAESTRO_PATH.stat()
    with open(CV_MAESTRO_PATH, "r", encoding="utf-8") as f:
        content = f.read()

    return {
        "exists": True,
        "filename": "cv-maestro.md",
        "content": content,
        "char_count": len(content),
        "updated_at": stat.st_mtime
    }
